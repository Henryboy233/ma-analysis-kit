#!/usr/bin/env python3
"""
生成多格式测试文件
用于测试商业尽调模块的文件解析功能
"""

import os
from datetime import datetime

# ========== 1. 生成 PDF 测试文件 ==========
def create_pdf():
    try:
        from fpdf import FPDF
        
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        # 标题
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(200, 10, txt="Test Company - Business Overview", ln=True, align='C')
        pdf.ln(10)
        
        # 内容
        pdf.set_font("Arial", size=12)
        content = """Company Name: Test Tech Co., Ltd.
Industry: Financial Technology
Founded: 2020
Employees: 150

Business Model:
We provide SaaS solutions for small and medium banks, helping them 
reduce operational costs through automated processing and AI-driven 
customer service systems.

Revenue Streams:
- Subscription fees (60%)
- Transaction commissions (30%)
- Consulting services (10%)

Key Partners:
- Major cloud service providers
- Payment gateway companies
- Regulatory technology firms

This is a test document for file parsing functionality."""
        
        for line in content.split('\n'):
            pdf.cell(200, 10, txt=line, ln=True)
        
        pdf.output("test_document.pdf")
        print("✅ PDF test file created: test_document.pdf")
        return True
    except ImportError:
        print("⚠️ fpdf not installed, skipping PDF generation")
        return False

# ========== 2. 生成 Word 测试文件 ==========
def create_word():
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        
        doc = Document()
        
        # 标题
        title = doc.add_heading('Test Company - Company Profile', 0)
        
        # 基本信息
        doc.add_heading('Basic Information', level=1)
        doc.add_paragraph('Company: Test Tech Solutions Inc.')
        doc.add_paragraph('Industry: Artificial Intelligence / Fintech')
        doc.add_paragraph('Stage: Series B')
        doc.add_paragraph('Revenue 2024: 50M CNY')
        
        # 核心业务
        doc.add_heading('Core Business', level=1)
        doc.add_paragraph(
            'We develop intelligent risk assessment systems for financial institutions. '
            'Our AI models analyze customer credit profiles in real-time.'
        )
        
        # 团队
        doc.add_heading('Management Team', level=1)
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Name'
        hdr_cells[1].text = 'Position'
        
        row1 = table.rows[1].cells
        row1[0].text = 'Zhang Wei'
        row1[1].text = 'CEO'
        
        row2 = table.rows[2].cells
        row2[0].text = 'Li Hua'
        row2[1].text = 'CTO'
        
        row3 = table.rows[3].cells
        row3[0].text = 'Wang Fang'
        row3[1].text = 'CFO'
        
        doc.add_paragraph()
        doc.add_paragraph('Test document for Word parsing.')
        
        doc.save("test_document.docx")
        print("✅ Word test file created: test_document.docx")
        return True
    except ImportError:
        print("⚠️ python-docx not installed, skipping Word generation")
        return False

# ========== 3. 生成 Excel 测试文件 ==========
def create_excel():
    try:
        import pandas as pd
        
        # 创建财务数据
        financial_data = {
            'Year': [2021, 2022, 2023, 2024],
            'Revenue(CNY)': [10000000, 25000000, 42000000, 68000000],
            'Cost(CNY)': [6000000, 15000000, 25200000, 38000000],
            'Gross_Profit(CNY)': [4000000, 10000000, 16800000, 30000000],
            'Net_Profit(CNY)': [500000, 2000000, 4200000, 8500000],
            'Employees': [30, 65, 110, 150]
        }
        
        # 创建业务指标数据
        metrics_data = {
            'Metric': ['Customer_Count', 'Retention_Rate', 'CAC', 'LTV', 'NPS'],
            '2023': [120, '75%', 5000, 35000, 45],
            '2024': [280, '82%', 4200, 42000, 52]
        }
        
        # 保存到不同sheet
        with pd.ExcelWriter('test_document.xlsx', engine='openpyxl') as writer:
            pd.DataFrame(financial_data).to_excel(writer, sheet_name='Financials', index=False)
            pd.DataFrame(metrics_data).to_excel(writer, sheet_name='Metrics', index=False)
            
            # 添加说明sheet
            notes = pd.DataFrame({
                'Note': ['This is a test Excel file for parsing', 
                        'Financial data is fictional',
                        'Created for M&A analysis system testing']
            })
            notes.to_excel(writer, sheet_name='Notes', index=False)
        
        print("✅ Excel test file created: test_document.xlsx")
        return True
    except ImportError:
        print("⚠️ pandas/openpyxl not installed, skipping Excel generation")
        return False

# ========== 4. 生成 PPT 测试文件 ==========
def create_ppt():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        
        prs = Presentation()
        
        # Slide 1: Title
        slide_layout = prs.slide_layouts[0]  # Title slide
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "Test Tech Inc."
        subtitle.text = "Series B Investment Pitch Deck\n2024"
        
        # Slide 2: Problem
        slide_layout = prs.slide_layouts[1]  # Title and content
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        title.text = "Problem Statement"
        content.text = """Traditional banks face challenges:
• High operational costs for customer onboarding
• Manual risk assessment takes 3-5 days
• Customer churn rate up to 40% annually
• Compliance costs increasing 20% per year"""
        
        # Slide 3: Solution
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        title.text = "Our Solution"
        content.text = """AI-Powered Banking Platform:
• Automated KYC in 2 minutes
• Real-time risk scoring
• Smart customer retention system
• Built-in compliance monitoring
• 60% cost reduction for clients"""
        
        # Slide 4: Market
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        title.text = "Market Opportunity"
        content.text = """• TAM: $50B (Global banking software)
• SAM: $8B (APAC region)
• SOM: $500M (Target segment)
• Growth rate: 25% CAGR
• 500+ potential enterprise clients"""
        
        # Slide 5: Funding
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        title.text = "Series B Investment"
        content.text = """Seeking: $30M Series B

Use of Funds:
• 40% R&D and product development
• 30% Market expansion
• 20% Key hires
• 10% Working capital

Contact: investor@testtech.com"""
        
        prs.save('test_presentation.pptx')
        print("✅ PowerPoint test file created: test_presentation.pptx")
        return True
    except ImportError:
        print("⚠️ python-pptx not installed, skipping PPT generation")
        return False

# ========== 生成简单的纯文本文件 ==========
def create_text():
    content = """Test Company - Executive Summary
================================

Company: Test Innovation Ltd.
Sector: E-commerce Technology
Stage: Growth

Overview:
This is a sample text document for testing the file upload 
and text extraction functionality of the M&A analysis system.

Key Points:
1. Revenue growth of 150% YoY
2. Customer base: 50,000+ active users
3. Key partnerships with major logistics providers
4. Technology stack: Cloud-native, microservices
5. Competitive advantage: AI-driven personalization

Contact: info@testinnovation.com

---
Test file generated on: {date}
""".format(date=datetime.now().strftime('%Y-%m-%d'))
    
    with open('test_readme.txt', 'w', encoding='utf-8') as f:
        f.write(content)
    print("✅ Text test file created: test_readme.txt")
    return True

# ========== 主程序 ==========
if __name__ == "__main__":
    print("=" * 50)
    print("Generating Test Files for M&A Analysis System")
    print("=" * 50)
    print()
    
    results = []
    results.append(("PDF", create_pdf()))
    results.append(("Word", create_word()))
    results.append(("Excel", create_excel()))
    results.append(("PowerPoint", create_ppt()))
    results.append(("Text", create_text()))
    
    print()
    print("=" * 50)
    print("Summary:")
    for name, success in results:
        status = "✅ Created" if success else "❌ Failed"
        print(f"  {name}: {status}")
    print("=" * 50)
